from __future__ import annotations

from pathlib import Path

from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw


OUT_DIR = Path(__file__).resolve().parent


def make_pdf() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.set_auto_page_break(auto=True, margin=15)
    lines = [
        "Certificate of Insurance",
        "Insured Entity: Acme Facilities LLC",
        "Insurer: Contoso Insurance Co.",
        "Coverage Amount: 1000000",
        "Expiry Date: 2027-01-31",
        "Additional Insured: Certa Corp",
    ]
    for line in lines:
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(OUT_DIR / "sample_coi.pdf"))


def make_docx() -> None:
    doc = Document()
    for line in [
        "Master Services Agreement",
        "Party A: Certa Corp",
        "Party B: Acme Facilities LLC",
        "Effective Date: 2026-01-01",
        "Term Length: 12 months",
        "Liability Cap: 500000",
    ]:
        doc.add_paragraph(line)
    doc.save(OUT_DIR / "sample_msa.docx")


def _draw_image(path: Path, title: str, lines: list[str], fmt: str) -> None:
    image = Image.new("RGB", (1200, 700), color="white")
    draw = ImageDraw.Draw(image)
    y = 40
    draw.text((40, y), title, fill="black")
    y += 60
    for line in lines:
        draw.text((40, y), line, fill="black")
        y += 40
    image.save(path, format=fmt)


def make_images() -> None:
    _draw_image(
        OUT_DIR / "sample_w9.png",
        "W-9 Form",
        [
            "Legal Entity Name: Acme Facilities LLC",
            "TIN: 12-3456789",
            "Entity Type: LLC",
            "Signature Date: 2025-05-01",
        ],
        "PNG",
    )
    _draw_image(
        OUT_DIR / "sample_bank.jpg",
        "Bank Details",
        [
            "Account Holder: Acme Facilities LLC",
            "Routing Number: 123456789",
            "Account Number: 9876543210",
            "Bank Name: Example National Bank",
        ],
        "JPEG",
    )


def make_email() -> None:
    (OUT_DIR / "sample_email.txt").write_text(
        "\n".join(
            [
                "Subject: Vendor onboarding submission",
                "",
                "Hello compliance team,",
                "Attached are the latest vendor documents for Acme Facilities LLC.",
                "Please let us know if you need anything else.",
            ]
        ),
        encoding="utf-8",
    )


def make_text_docs() -> None:
    (OUT_DIR / "sample_w9.txt").write_text(
        "\n".join(
            [
                "W-9 Form",
                "Legal Entity Name: Nexus Data Solutions LLC",
                "TIN: 47-3821956",
                "Entity Type: LLC",
                "Signature Date: 2025-05-01",
            ]
        ),
        encoding="utf-8",
    )
    (OUT_DIR / "sample_bank.txt").write_text(
        "\n".join(
            [
                "Bank Details",
                "Account Holder: Nexus Data Solutions LLC",
                "Routing Number: 123456789",
                "Account Number: 9876543210",
                "Bank Name: Example National Bank",
            ]
        ),
        encoding="utf-8",
    )
    (OUT_DIR / "sample_dpa.txt").write_text(
        "\n".join(
            [
                "Data Processing Agreement",
                "Data Categories: employee records",
                "Retention Period: 2 years",
                "Sub-processor: CloudHost One",
                "Breach Notification Hours: 48",
            ]
        ),
        encoding="utf-8",
    )
    (OUT_DIR / "sample_soc2.txt").write_text(
        "\n".join(
            [
                "SOC 2 Type II Report",
                "Period Start: 2025-01-01",
                "Period End: 2025-12-31",
                "Covered Services: managed analytics platform",
                "Auditor: Assurance Partners LLP",
            ]
        ),
        encoding="utf-8",
    )


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_pdf()
    make_docx()
    make_images()
    make_email()
    make_text_docs()
    print(f"Generated sample docs in {OUT_DIR}")
